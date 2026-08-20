"""Importacion y exportacion de estudiantes en formato Excel (.xlsx).

Usa openpyxl (sin dependencias binarias, compatible con PyInstaller) para
leer y escribir hojas de calculo. La logica de validacion reutiliza las
mismas reglas que el formulario de matricula en `main.py`.
"""

from openpyxl import Workbook, load_workbook

import database

COLUMNAS_EXPORT = ("Codigo_Matricula", "DNI", "Nombres", "Apellidos", "Telefono", "Aula")
COLUMNAS_IMPORT_REQUERIDAS = ("DNI", "Nombres", "Apellidos", "Telefono")


def exportar_alumnos(ruta):
    """Escribe todos los estudiantes registrados en un archivo .xlsx."""
    alumnos = database.obtener_alumnos()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Estudiantes"
    hoja.append(COLUMNAS_EXPORT)
    for codigo, dni, nombres, apellidos, telefono, _id_aula, nombre_aula in alumnos:
        hoja.append([codigo, dni, nombres, apellidos, telefono or "", nombre_aula or ""])

    for columna, ancho in zip("ABCDEF", (16, 12, 22, 22, 14, 26)):
        hoja.column_dimensions[columna].width = ancho

    libro.save(ruta)
    return True, f"{len(alumnos)} estudiante(s) exportado(s) correctamente."


def importar_alumnos(ruta, id_aula_defecto, usuario):
    """Lee un .xlsx y registra los estudiantes validos.

    Devuelve (cantidad_insertada, lista_de_errores). Cada fila se valida con
    las mismas reglas del formulario de matricula (DNI de 8 digitos, celular
    de 9, nombres/apellidos obligatorios); las filas invalidas se reportan sin
    detener la importacion del resto, y cada mensaje indica que se esperaba
    para que el archivo se pueda corregir sin adivinar. Los alumnos validos se
    insertan en una sola conexion/transaccion (antes se abria una conexion
    SQLite por cada fila del archivo).
    """
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = libro.active
        filas = list(hoja.iter_rows(values_only=True))
    finally:
        libro.close()
    if not filas:
        return 0, ["El archivo esta vacio."]

    encabezado = [str(valor).strip() if valor is not None else "" for valor in filas[0]]
    indices = {nombre: posicion for posicion, nombre in enumerate(encabezado) if nombre}

    faltantes = [col for col in COLUMNAS_IMPORT_REQUERIDAS if col not in indices]
    if faltantes:
        return 0, [
            f"Faltan columnas obligatorias: {', '.join(faltantes)}. "
            f"La primera fila del Excel debe tener los encabezados: {', '.join(COLUMNAS_IMPORT_REQUERIDAS)} "
            "(opcionales: Codigo_Matricula y Aula)."
        ]

    aulas_por_nombre = {nombre: id_aula for id_aula, nombre in database.obtener_aulas_activas()}
    aulas_disponibles = ", ".join(aulas_por_nombre) if aulas_por_nombre else "no hay aulas activas creadas"

    def obtener(fila, columna):
        indice = indices.get(columna)
        if indice is None or indice >= len(fila):
            return ""
        valor = fila[indice]
        return str(valor).strip() if valor is not None else ""

    filas_validas = []
    errores = []

    for numero_fila, fila in enumerate(filas[1:], start=2):
        if fila is None or all(valor in (None, "") for valor in fila):
            continue

        dni = obtener(fila, "DNI")
        nombres = obtener(fila, "Nombres").upper()
        apellidos = obtener(fila, "Apellidos").upper()
        telefono = obtener(fila, "Telefono")
        codigo = obtener(fila, "Codigo_Matricula").upper()
        nombre_aula = obtener(fila, "Aula")

        if not dni.isdigit() or len(dni) != 8:
            errores.append(
                f"Fila {numero_fila}: DNI invalido ('{dni or 'vacio'}'). Debe tener exactamente 8 digitos numericos."
            )
            continue
        if not nombres or not apellidos:
            faltante = "Nombres" if not nombres else "Apellidos"
            errores.append(f"Fila {numero_fila}: falta el dato de '{faltante}'.")
            continue
        if not telefono.isdigit() or len(telefono) != 9:
            errores.append(
                f"Fila {numero_fila}: celular invalido ('{telefono or 'vacio'}'). Debe tener exactamente 9 digitos numericos."
            )
            continue

        if nombre_aula:
            id_aula = aulas_por_nombre.get(nombre_aula)
            if id_aula is None:
                errores.append(
                    f"Fila {numero_fila}: el aula '{nombre_aula}' no existe o no esta activa. "
                    f"Aulas disponibles: {aulas_disponibles}."
                )
                continue
        else:
            id_aula = id_aula_defecto

        if id_aula is None:
            errores.append(
                f"Fila {numero_fila}: no se indico aula y no hay aula por defecto seleccionada. "
                f"Elija una en 'Aula por defecto', o agregue la columna 'Aula' con uno de estos valores: "
                f"{aulas_disponibles}."
            )
            continue

        filas_validas.append([numero_fila, dni, nombres, apellidos, telefono, codigo, id_aula])

    filas_sin_codigo = [fila_valida for fila_valida in filas_validas if not fila_valida[5]]
    if filas_sin_codigo:
        codigos_nuevos = database.obtener_codigos_matricula_disponibles(len(filas_sin_codigo))
        for fila_valida, codigo_nuevo in zip(filas_sin_codigo, codigos_nuevos):
            fila_valida[5] = codigo_nuevo

    if not filas_validas:
        return 0, errores

    registros = [
        (fila_valida[5], fila_valida[1], fila_valida[2], fila_valida[3], fila_valida[4], fila_valida[6])
        for fila_valida in filas_validas
    ]
    resultados = database.insertar_alumnos_lote(registros, usuario)

    insertados = 0
    for fila_valida, (exito, mensaje) in zip(filas_validas, resultados):
        if exito:
            insertados += 1
        else:
            numero_fila, dni = fila_valida[0], fila_valida[1]
            errores.append(f"Fila {numero_fila} (DNI {dni}): {mensaje}")

    return insertados, errores
